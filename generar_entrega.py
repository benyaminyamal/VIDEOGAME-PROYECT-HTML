#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el documento de entrega (.docx) para la actividad sumativa:
"Optimizando la lógica y rendimiento de una página web con JavaScript".
Incluye portada, enlaces y espacios rotulados para pegar las capturas.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MORADO = RGBColor(0x7C, 0x3A, 0xED)
GRIS = RGBColor(0x55, 0x55, 0x55)

REPO_URL = "https://github.com/benyaminyamal/VIDEOGAME-PROYECT-HTML"
DEPLOY_URL = "https://benyaminyamal.github.io/VIDEOGAME-PROYECT-HTML/"

doc = Document()

# Márgenes
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)

# Fuente base
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(11)


def add_placeholder(texto):
    """Agrega un recuadro con borde punteado que indica dónde pegar la captura."""
    tabla = doc.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = tabla.cell(0, 0)
    celda.width = Inches(6.3)
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n[  Pega aquí tu captura de pantalla  ]\n" + texto + "\n\n")
    run.italic = True
    run.font.color.rgb = GRIS
    run.font.size = Pt(10)
    # Bordes de la celda
    tcPr = celda._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "7C3AED")
        borders.append(el)
    tcPr.append(borders)
    doc.add_paragraph()


def titulo(texto, size=15):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = MORADO
    return p


# ------------------- PORTADA -------------------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TodoJuegos")
r.bold = True
r.font.size = Pt(30)
r.font.color.rgb = MORADO

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Actividad Sumativa – Semana 6\n"
              "Optimizando la lógica y rendimiento de una página web con JavaScript")
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("eCommerce de videojuegos con Bootstrap 5, JavaScript y Fetch API")
r.italic = True
r.font.color.rgb = GRIS

# Datos del estudiante (para completar)
doc.add_paragraph()
tbl = doc.add_table(rows=3, cols=2)
tbl.style = "Light List Accent 1"
datos = [("Nombre del estudiante:", ""), ("Asignatura:", ""), ("Fecha de entrega:", "")]
for i, (k, v) in enumerate(datos):
    tbl.cell(i, 0).paragraphs[0].add_run(k).bold = True
    tbl.cell(i, 1).text = v

doc.add_paragraph()

# ------------------- ENLACES -------------------
titulo("1. Enlaces del proyecto")
p = doc.add_paragraph()
p.add_run("Repositorio en GitHub (público): ").bold = True
p.add_run(REPO_URL)
p = doc.add_paragraph()
p.add_run("Despliegue público (GitHub Pages): ").bold = True
p.add_run(DEPLOY_URL)
p = doc.add_paragraph()
r = p.add_run("Nota: la rama gh-pages está configurada para publicar el sitio automáticamente.")
r.italic = True
r.font.color.rgb = GRIS

doc.add_page_break()

# ------------------- CAPTURAS -------------------
titulo("2. Capturas de pantalla")

# 2.1 Estructura
titulo("2.1  Estructura de la página", 13)
doc.add_paragraph(
    "Vista general del sitio: barra de navegación (categorías + buscador + carrito), "
    "sección \u201cProductos destacados\u201d con las tarjetas de videojuegos, y pie de página "
    "con datos de contacto y redes sociales.")
add_placeholder("Barra de navegación + productos")
add_placeholder("Pie de página (contacto y redes sociales)")

doc.add_page_break()

# 2.2 Interacciones dinámicas
titulo("2.2  Interacciones dinámicas con JavaScript", 13)

doc.add_paragraph().add_run("a) Agregar productos al carrito (evento click)").bold = True
doc.add_paragraph(
    "Se hace clic en el botón \u201cAgregar\u201d de varios productos y se muestra el panel "
    "del carrito con el resumen (productos, cantidades y total) y el contador actualizado.")
add_placeholder("Carrito con productos agregados y total")

doc.add_paragraph().add_run("b) Búsqueda de productos (evento submit)").bold = True
doc.add_paragraph(
    "Se escribe un término en el buscador (por ejemplo \u201celden\u201d o \u201cgod\u201d) y se envía el "
    "formulario; la lista de productos se filtra dinámicamente mostrando solo coincidencias.")
add_placeholder("Resultado del buscador filtrado")

doc.add_paragraph().add_run("c) Filtro por categoría (opcional)").bold = True
doc.add_paragraph(
    "Al hacer clic en una categoría del menú (Acción, RPG o Aventura) se muestran solo los "
    "juegos de esa categoría.")
add_placeholder("Productos filtrados por categoría")

doc.add_page_break()

# 2.3 Fetch API
titulo("2.3  Carga de datos externos con la Fetch API", 13)
doc.add_paragraph(
    "Con las herramientas de desarrollador del navegador (F12) → pestaña \u201cNetwork/Red\u201d → "
    "filtro \u201cFetch/XHR\u201d, se recarga la página y se observa la petición al archivo "
    "productos.json con estado 200 y su respuesta (JSON con la lista de videojuegos).")
add_placeholder("Petición a productos.json (Network) + respuesta JSON")

doc.add_paragraph().add_run("Gestión de errores (opcional)").bold = True
doc.add_paragraph(
    "Simulando el modo \u201cOffline\u201d en la pestaña Network y recargando, la página muestra un "
    "mensaje amigable de error en lugar de fallar silenciosamente.")
add_placeholder("Mensaje amigable de error")

doc.add_page_break()

# ------------------- RESUMEN TÉCNICO -------------------
titulo("3. Resumen de cumplimiento de requisitos")
requisitos = [
    ("Bootstrap 5", "Navbar responsiva con 2+ categorías, grid de productos y footer; diseño accesible y adaptable a móviles."),
    ("Evento click", "Botón \u201cAgregar\u201d añade productos al carrito."),
    ("Evento submit", "Formulario de búsqueda procesa y filtra los productos."),
    ("Manipulación del DOM", "El carrito se genera dinámicamente con su resumen y total."),
    ("Fetch API", "La lista de productos se carga desde productos.json (archivo JSON local)."),
    ("Gestión de errores", "Mensaje amigable cuando los datos no cargan correctamente."),
    ("Buenas prácticas", "Código en funciones claras y reutilizables, con comentarios explicativos."),
]
tabla = doc.add_table(rows=1, cols=2)
tabla.style = "Light Grid Accent 1"
hdr = tabla.rows[0].cells
hdr[0].paragraphs[0].add_run("Requisito").bold = True
hdr[1].paragraphs[0].add_run("Cómo se cumple").bold = True
for req, desc in requisitos:
    fila = tabla.add_row().cells
    fila[0].paragraphs[0].add_run(req).bold = True
    fila[1].text = desc

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Estructura de archivos: index.html, inicio.html, estilos.css, app.js, productos.json")
r.italic = True
r.font.color.rgb = GRIS

salida = "/Users/benjaminclaros/Projects/VIDEOGAME-PROYECT-HTML/Entrega_TodoJuegos_S6.docx"
doc.save(salida)
print("Documento generado:", salida)
